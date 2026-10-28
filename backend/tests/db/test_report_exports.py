"""Report exports and report settings (WP-035 session 3). Test names cite parity file 09.

The incident world of test_reports_api.py (the real incident source); object storage is
an in-memory fake (nothing touches the dev stack's MinIO); the worker's subject loader is
identity_access.subject_for, as in the consumers worker (lane B, #82).
"""

import csv
import io
import uuid
from collections.abc import AsyncIterator
from datetime import UTC, datetime, timedelta
from typing import Any

import pytest
import pytest_asyncio
from openpyxl import load_workbook

from app.core.database import Database, create_engine
from app.core.errors import ServiceUnavailable
from app.core.tenancy import TenantContext
from app.modules.analytics.api.routes import STORAGE_STATE
from app.modules.analytics.service.exports import (
    SubjectLoader,
    Worker,
    build_export,
    fail_stalled_exports,
)
from app.modules.identity_access import subject_for
from app.shared.labels import label
from tests.db import test_reports_api
from tests.db.test_custom_fields_api import add_field, choices
from tests.db.test_incidents_api import CONSOLE, Org, World
from tooling.db_testing import ThrowawayDatabase

pytestmark = pytest.mark.db
world = test_reports_api.world
reports = test_reports_api.reports

OPEN = {"field": "status_category", "op": "in", "value": ["open", "paused"]}


class MemoryStorage:
    """ObjectStorage in a dict (the port's methods the export path uses)."""

    def __init__(self) -> None:
        self.files: dict[str, bytes] = {}

    async def put(self, key: str, data: bytes, content_type: str) -> None:
        del content_type
        self.files[key] = data

    async def put_stream(self, key: str, parts: AsyncIterator[bytes], content_type: str) -> int:
        del content_type
        data = b"".join([part async for part in parts])
        self.files[key] = data
        return len(data)

    async def get(self, key: str) -> bytes | None:
        return self.files.get(key)

    async def delete(self, key: str) -> None:
        self.files.pop(key, None)

    async def check(self) -> None:
        return None

    async def read_range(self, key: str, start: int, end: int) -> AsyncIterator[bytes]:
        yield self.files[key][start : end + 1]


@pytest_asyncio.fixture
async def exporting(reports: World) -> AsyncIterator[tuple[World, MemoryStorage]]:
    storage = MemoryStorage()
    setattr(reports.app.state, STORAGE_STATE, storage)
    try:
        yield reports, storage
    finally:
        setattr(reports.app.state, STORAGE_STATE, None)


async def _request(world: World, org: Any, user: uuid.UUID, **body: Any) -> Any:  # noqa: ANN401 - JSON
    return await world.call(
        "POST",
        "/reports/exports",
        org,
        user,
        json={"title": "Open incidents", "definition": {"entity": "incident"}, **body},
        headers=CONSOLE,
    )


async def _build(
    migrated_db: ThrowawayDatabase,
    storage: MemoryStorage,
    org: Org,
    export_id: str,
    load_subject: SubjectLoader = subject_for,
) -> None:
    """What the consumers worker does for one requested export."""
    database = Database(
        create_engine(migrated_db.app_url, application_name="t", use_null_pool=True)
    )
    try:
        async with database.tenant_session(TenantContext(org.tenant)) as session:
            await build_export(
                session,
                uuid.UUID(export_id),
                Worker(
                    reporting=database,
                    storage=storage,
                    load_subject=load_subject,
                    render_pdf=None,
                    timeout_seconds=30,
                    max_rows=1000,
                    link_days=7,
                ),
            )
    finally:
        await database.dispose()


async def test_req_rpt_058_records_csv_holds_exactly_what_the_user_may_list(
    exporting: tuple[World, MemoryStorage], migrated_db: ThrowawayDatabase
) -> None:
    world, storage = exporting
    a = world.a
    await test_reports_api.seed(world)
    await world.create(a, a.admin, group_id=str(a.ga), subject='=HYPERLINK("http://x","y")')
    requested = await _request(
        world,
        a,
        a.tech_a,
        format="csv",
        content="records",
        definition={"entity": "incident", "filter": OPEN, "group_by": [{"field": "priority"}]},
    )
    assert requested.status_code == 202, requested.text
    export = requested.json()
    assert export["status"] == "running"
    await _build(migrated_db, storage, a, export["id"])

    listed = await world.call("GET", "/reports/exports", a, a.tech_a)
    mine = next(e for e in listed.json() if e["id"] == export["id"])
    assert mine["status"] == "ready", mine
    visible = await test_reports_api.listed(world, a, a.tech_a, OPEN)
    assert mine["row_count"] == len(visible)  # tech_a's scope (D-052), nothing more

    download = await world.call("GET", f"/reports/exports/{export['id']}/file", a, a.tech_a)
    assert download.status_code == 200, download.text
    assert download.content.startswith("﻿".encode())
    rows = list(csv.reader(io.StringIO(download.content.decode("utf-8-sig"))))
    table = rows[rows.index(next(r for r in rows if r and r[0] == "Priority")) :]
    ids = {r[1] for r in table[1:]}
    assert len(ids) == len(visible)
    subjects = [r[2] for r in table[1:]]
    assert '\'=HYPERLINK("http://x","y")' in subjects  # neutralised (D-100)
    assert not any(s.startswith("=") for s in subjects)
    assert not any(r[0] in {"p1", "p2", "p3", "p4"} for r in table[1:])  # labels, never codes
    header = table[0]
    groups = {r[header.index("Group")] for r in table[1:]}
    assert "Alpha GA" in groups  # named by organization's resolver, not an id
    statuses = {r[header.index("Status")] for r in table[1:]}
    assert statuses <= {"New", "Not set"}, statuses
    assert "Unknown value" not in {cell for r in table[1:] for cell in r}

    # Only the person who asked can download it.
    other = await world.call("GET", f"/reports/exports/{export['id']}/file", a, a.tech_b)
    assert other.status_code == 404
    audit = await world.sql(
        "SELECT action, details FROM audit_entries WHERE tenant_id = $1 AND action LIKE 'report.%'",
        a.tenant,
    )
    actions = {r["action"] for r in audit}
    assert {"report.exported", "report.export_downloaded"} <= actions
    exported = next(r for r in audit if r["action"] == "report.exported")
    assert '"rows": ' + str(len(visible)) in str(exported["details"]).replace("'", '"')


async def test_req_rpt_058_ac1_xlsx_summary_and_records_with_totals(
    exporting: tuple[World, MemoryStorage], migrated_db: ThrowawayDatabase
) -> None:
    world, storage = exporting
    a = world.a
    await test_reports_api.seed(world)
    requested = await _request(
        world,
        a,
        a.lead,
        format="xlsx",
        content="records",
        language="ar",
        definition={"entity": "incident", "group_by": [{"field": "status_category"}]},
    )
    assert requested.status_code == 202, requested.text
    await _build(migrated_db, storage, a, requested.json()["id"])
    data = (
        await world.call("GET", f"/reports/exports/{requested.json()['id']}/file", a, a.lead)
    ).content
    workbook = load_workbook(io.BytesIO(data))
    summary, records = workbook.worksheets
    assert summary.sheet_view.rightToLeft is True
    values = [[c.value for c in row] for row in summary.iter_rows()]
    total = values[-1]
    assert total[1] == 6  # every incident of org A, a number not text
    assert isinstance(total[1], int)
    assert records.max_row - 6 == 6  # header lines + blank + header, then six records


async def test_exports_need_permission_and_storage(exporting: tuple[World, MemoryStorage]) -> None:
    world, _storage = exporting
    a = world.a
    refused = await _request(world, a, a.requester, format="csv")
    assert refused.status_code == 403
    setattr(world.app.state, STORAGE_STATE, None)
    no_storage = await _request(world, a, a.tech_a, format="csv")
    assert no_storage.status_code == 503
    assert no_storage.json()["code"] == "storage_not_configured"


async def test_pdf_without_the_worker_image_fails_visibly(
    exporting: tuple[World, MemoryStorage], migrated_db: ThrowawayDatabase
) -> None:
    world, storage = exporting
    a = world.a
    requested = await _request(world, a, a.tech_a, format="pdf")
    await _build(migrated_db, storage, a, requested.json()["id"])
    mine = (await world.call("GET", "/reports/exports", a, a.tech_a)).json()[0]
    assert (mine["status"], mine["problem"]) == ("failed", "pdf_not_available")
    events = await world.sql(
        "SELECT id FROM outbox_events WHERE tenant_id = $1"
        " AND event_type = 'analytics.report_export.finished'",
        a.tenant,
    )
    assert events  # the person who asked is told live (REQ-RPT-056)


async def test_d041_report_settings_are_editable_by_admins_and_shape_weeks(reports: World) -> None:
    world = reports
    a = world.a
    got = await world.call("GET", "/reports/settings", a, a.tech_a)
    assert got.status_code == 200, got.text
    assert got.json()["week_start"] == "monday"
    assert got.json()["interactive_timeout_seconds"] == 30
    version = got.json()["version"]
    refused = await world.call(
        "PATCH",
        "/reports/settings",
        a,
        a.tech_a,
        json={"week_start": "sunday"},
        headers={**CONSOLE, "If-Match": f'"{version}"'},
    )
    assert refused.status_code == 403
    changed = await world.call(
        "PATCH",
        "/reports/settings",
        a,
        a.admin,
        json={"week_start": "sunday", "export_link_days": 3},
        headers={**CONSOLE, "If-Match": f'"{version}"'},
    )
    assert changed.status_code == 200, changed.text
    assert changed.json()["week_start"] == "sunday"
    stale = await world.call(
        "PATCH",
        "/reports/settings",
        a,
        a.admin,
        json={"export_link_days": 5},
        headers={**CONSOLE, "If-Match": f'"{version}"'},
    )
    assert stale.status_code == 409
    out_of_range = await world.call(
        "PATCH",
        "/reports/settings",
        a,
        a.admin,
        json={"interactive_timeout_seconds": 600},
        headers={**CONSOLE, "If-Match": f'"{changed.json()["version"]}"'},
    )
    assert out_of_range.status_code == 422

    # A Wednesday incident counts in the week that starts on the Sunday before it.
    item = await world.create(a, a.admin)
    await world.sql(
        "UPDATE work_items SET created_at = '2026-10-21T10:00:00+00' WHERE id = $1",
        uuid.UUID(item["id"]),
    )
    weeks = await test_reports_api.ran(
        world,
        a,
        a.admin,
        filter={
            "field": "created_at",
            "op": "between",
            "value": ["2026-10-18T00:00:00Z", "2026-10-25T00:00:00Z"],
        },
        group_by=[{"field": "created_at", "bucket": "week"}],
    )
    assert [r["keys"][0][:10] for r in weeks["rows"]] == ["2026-10-18"]  # a Sunday
    audit = await world.sql(
        "SELECT outcome FROM audit_entries WHERE tenant_id = $1"
        " AND action = 'report.settings_changed'",
        a.tenant,
    )
    # The change, and the three refusals (permission, stale version, out of range), are audited.
    assert sorted(r["outcome"] for r in audit).count("done") == 1
    assert len(audit) == 4


async def test_categories_are_named_by_their_full_path_like_the_screens(
    exporting: tuple[World, MemoryStorage], migrated_db: ThrowawayDatabase
) -> None:
    world, storage = exporting
    a = world.a
    await world.call("GET", "/incidents/categories", a, a.admin)  # seeds the tenant's tree
    printer = await world.sql(
        "SELECT id FROM work_item_categories WHERE tenant_id = $1 AND code = 'printer'", a.tenant
    )
    assert printer, "the default category tree has a printer under hardware"
    await world.create(a, a.admin, category_id=str(printer[0]["id"]))
    requested = await _request(
        world,
        a,
        a.admin,
        format="csv",
        definition={"entity": "incident", "group_by": [{"field": "category_id"}]},
    )
    await _build(migrated_db, storage, a, requested.json()["id"])
    data = (
        await world.call("GET", f"/reports/exports/{requested.json()['id']}/file", a, a.admin)
    ).content.decode("utf-8-sig")
    assert "Hardware \u203a Printer" in data  # "Hardware > Printer" with the angle quote


async def no_entitlements(session: Any, tenant_id: uuid.UUID, user_id: uuid.UUID) -> Any:  # noqa: ANN401 - SubjectLoader
    """A worker without the entitlement service, as the consumers worker was (e2e, 2026-10-28)."""
    del session, tenant_id, user_id
    raise ServiceUnavailable(code="entitlements_unavailable")


async def test_req_rpt_056_a_worker_missing_a_service_fails_the_export_visibly(
    exporting: tuple[World, MemoryStorage], migrated_db: ThrowawayDatabase
) -> None:
    world, storage = exporting
    a = world.a
    export = (await _request(world, a, a.tech_a, format="csv")).json()
    await _build(migrated_db, storage, a, export["id"], load_subject=no_entitlements)
    mine = next(
        e
        for e in (await world.call("GET", "/reports/exports", a, a.tech_a)).json()
        if e["id"] == export["id"]
    )
    assert (mine["status"], mine["problem"]) == ("failed", "report_failed")  # not "running"
    assert storage.files == {}


async def test_req_rpt_056_an_export_lost_by_its_worker_is_failed_by_the_sweep(
    exporting: tuple[World, MemoryStorage], migrated_db: ThrowawayDatabase
) -> None:
    world, _storage = exporting
    a = world.a
    lost = (await _request(world, a, a.tech_a, format="csv")).json()
    database = Database(
        create_engine(migrated_db.app_url, application_name="t", use_null_pool=True)
    )
    try:
        now = datetime.now(UTC)
        async with database.tenant_session(TenantContext(a.tenant)) as session:
            assert await fail_stalled_exports(session, now - timedelta(hours=1)) == 0  # recent
        async with database.tenant_session(TenantContext(a.tenant)) as session:
            assert await fail_stalled_exports(session, now + timedelta(seconds=1)) == 1
    finally:
        await database.dispose()
    mine = next(
        e
        for e in (await world.call("GET", "/reports/exports", a, a.tech_a)).json()
        if e["id"] == lost["id"]
    )
    assert (mine["status"], mine["problem"]) == ("failed", "report_failed")
    finished = await world.sql(
        "SELECT 1 FROM outbox_events WHERE tenant_id = $1 AND aggregate_id = $2"
        " AND event_type = 'analytics.report_export.finished'",
        a.tenant,
        lost["id"],
    )
    assert finished  # the person sees it fail live


@pytest.mark.parametrize("language", ["en", "ar"])
async def test_bug_01_no_export_header_reads_unknown_value(
    exporting: tuple[World, MemoryStorage], migrated_db: ThrowawayDatabase, language: str
) -> None:
    """Records exports and a grouping by a custom field name every column (QA WP-035 BUG-01)."""
    world, storage = exporting
    a = world.a
    await add_field(
        world,
        a,
        field_type="choice",
        labels={"en": "Floor", "ar": "الطابق"},
        choices=choices("Ground", "First"),
    )
    await world.create(a, a.admin, custom={"floor": "first"})
    requested = await _request(
        world,
        a,
        a.admin,
        format="xlsx",
        content="records",
        language=language,
        definition={"entity": "incident", "group_by": [{"field": "custom.floor"}]},
    )
    assert requested.status_code == 202, requested.text
    await _build(migrated_db, storage, a, requested.json()["id"])
    data = (
        await world.call("GET", f"/reports/exports/{requested.json()['id']}/file", a, a.admin)
    ).content
    unknown = label("reports.unknown_value", language)
    for sheet in load_workbook(io.BytesIO(data)).worksheets:
        cells = {c.value for row in sheet.iter_rows() for c in row}
        assert unknown not in cells, sheet.title
    summary, records = load_workbook(io.BytesIO(data)).worksheets
    floor = "الطابق" if language == "ar" else "Floor"
    assert floor in {c.value for row in summary.iter_rows() for c in row}
    header = next([c.value for c in row] for row in records.iter_rows() if row[0].value == floor)
    assert header[1:3] == (["المعرّف", "الموضوع"] if language == "ar" else ["ID", "Subject"])


async def test_records_export_takes_chosen_columns_and_sort_for_list_exports(
    exporting: tuple[World, MemoryStorage], migrated_db: ThrowawayDatabase
) -> None:
    """The incident list's export (WP-020d) reuses this job: its columns, its sort, its scope."""
    world, storage = exporting
    a = world.a
    await test_reports_api.seed(world)
    await world.create(a, a.admin, subject="Aardvark first", group_id=str(a.ga))
    requested = await _request(
        world,
        a,
        a.tech_a,
        format="csv",
        content="records",
        definition={
            "entity": "incident",
            "columns": ["subject", "display_id", "group_id"],
            "sort": [{"field": "subject", "desc": False}],
        },
    )
    assert requested.status_code == 202, requested.text
    await _build(migrated_db, storage, a, requested.json()["id"])
    data = (
        await world.call("GET", f"/reports/exports/{requested.json()['id']}/file", a, a.tech_a)
    ).content
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    header_at = next(i for i, r in enumerate(rows) if r[:2] == ["Subject", "ID"])
    records = rows[header_at + 1 :]
    assert rows[header_at] == ["Subject", "ID", "Group"]
    assert records[0][0] == "Aardvark first"  # sorted by subject
    visible = await test_reports_api.listed(world, a, a.tech_a, None)
    assert len(records) == len(visible)  # the requester's scope, nothing more


@pytest.mark.parametrize(
    ("definition", "field", "code"),
    [
        ({"columns": ["subject", "nope"]}, "columns[1]", "unknown_column"),
        ({"columns": ["subject", "subject"]}, "columns", "duplicate_column"),
        ({"sort": [{"field": "nope"}]}, "sort[0].field", "unknown_column"),
    ],
)
async def test_chosen_columns_and_sort_are_checked(
    exporting: tuple[World, MemoryStorage],
    definition: dict[str, Any],
    field: str,
    code: str,
) -> None:
    world, _storage = exporting
    a = world.a
    refused = await _request(
        world,
        a,
        a.admin,
        format="csv",
        content="records",
        definition={"entity": "incident", **definition},
    )
    assert refused.status_code == 422
    assert {(e["field"], e["code"]) for e in refused.json()["errors"]} == {(field, code)}
