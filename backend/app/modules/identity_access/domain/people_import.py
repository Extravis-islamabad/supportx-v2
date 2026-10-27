"""Reading a people file, CSV or Excel (WP-021, REQ-ADM-003). No database here.

Limits protect the import job from oversized and "bomb" files (product owner, WP-021): the file
size, the rows, the columns and the characters per cell are capped, and an Excel file (a zip) is
checked before it is opened: entry count, total unpacked size and compression ratio. Excel is
read with openpyxl in read-only mode, values only (formulas are never evaluated).
"""

import csv
import io
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass
from enum import StrEnum

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 40
MAX_CELL_CHARS = 5000
MAX_ZIP_ENTRIES = 200
MAX_UNPACKED_BYTES = 50 * 1024 * 1024
MAX_RATIO = 100  # unpacked / packed, per entry
RATIO_FROM_BYTES = 1024  # tiny entries (empty parts) compress far better and are harmless

# Column heading (lowercase, spaces or underscores) → field. Field codes and English headings.
COLUMNS: dict[str, str] = {
    "first name": "first_name",
    "last name": "last_name",
    "display name": "display_name",
    "email": "email",
    "email address": "email",
    "employee id": "employee_id",
    "job title": "job_title",
    "phone": "phone",
    "mobile": "mobile",
    "site": "site",
    "department": "department",
    "manager": "manager",
    "manager email": "manager",
    "type": "user_type",
    "user type": "user_type",
    "vip": "vip",
    "can sign in": "can_sign_in",
    "sign in": "can_sign_in",
    "language": "locale",
    "locale": "locale",
}
TRUE_WORDS = frozenset({"yes", "y", "true", "1", "x"})
FALSE_WORDS = frozenset({"no", "n", "false", "0", ""})


class ImportFailure(StrEnum):
    """Why a whole file was refused (labels `people_import_failure.<code>`)."""

    TOO_LARGE = "file_too_large"
    NOT_READABLE = "file_not_readable"
    NOT_UTF8 = "file_not_utf8"
    UNSAFE = "file_unsafe"
    TOO_MANY_ROWS = "too_many_rows"
    TOO_MANY_COLUMNS = "too_many_columns"
    CELL_TOO_LONG = "cell_too_long"
    NO_HEADER = "no_header"
    UNKNOWN_COLUMN = "unknown_column"
    NO_NAME_COLUMN = "no_name_column"
    EMPTY = "file_empty"


class ImportRefused(Exception):
    """The file cannot be imported at all."""

    def __init__(self, failure: ImportFailure, detail: str | None = None) -> None:
        """`detail`: e.g. the unknown column's heading (shown to the admin)."""
        super().__init__(failure.value)
        self.failure = failure
        self.detail = detail


@dataclass(frozen=True)
class SheetRows:
    """The mapped header and the data rows (row numbers as in the file, header = 1)."""

    fields: list[str]
    rows: list[tuple[int, dict[str, str]]]


def file_format(file_name: str, content: bytes) -> str:
    """The file's kind, csv or xlsx: by the content (an xlsx is a zip) and the extension."""
    name = file_name.lower()
    if content.startswith(b"PK\x03\x04") and name.endswith(".xlsx"):
        return "xlsx"
    if name.endswith((".csv", ".txt")) and not content.startswith(b"PK\x03\x04"):
        return "csv"
    raise ImportRefused(ImportFailure.NOT_READABLE)


def flag(value: str) -> bool | None:
    """Yes/no cells; None when the cell says neither."""
    word = value.strip().lower()
    if word in TRUE_WORDS:
        return True
    if word in FALSE_WORDS:
        return False
    return None


def read_rows(file_format_code: str, content: bytes) -> SheetRows:
    """Every data row as {field: text}; refuses the whole file when a limit is crossed."""
    if len(content) > MAX_FILE_BYTES:
        raise ImportRefused(ImportFailure.TOO_LARGE)
    raw = _xlsx_rows(content) if file_format_code == "xlsx" else _csv_rows(content)
    header: list[str] | None = None
    rows: list[tuple[int, dict[str, str]]] = []
    for number, cells in enumerate(raw, start=1):
        if len(cells) > MAX_COLUMNS:
            raise ImportRefused(ImportFailure.TOO_MANY_COLUMNS)
        if any(len(cell) > MAX_CELL_CHARS for cell in cells):
            raise ImportRefused(ImportFailure.CELL_TOO_LONG, str(number))
        if header is None:
            header = _fields(cells)
            continue
        if not any(cell.strip() for cell in cells):
            continue
        if len(rows) >= MAX_ROWS:
            raise ImportRefused(ImportFailure.TOO_MANY_ROWS)
        values = {
            field: cells[index].strip()
            for index, field in enumerate(header)
            if field and index < len(cells)
        }
        rows.append((number, values))
    if header is None:
        raise ImportRefused(ImportFailure.NO_HEADER)
    if not rows:
        raise ImportRefused(ImportFailure.EMPTY)
    return SheetRows([f for f in header if f], rows)


def _fields(cells: list[str]) -> list[str]:
    fields: list[str] = []
    for cell in cells:
        heading = " ".join(cell.strip().lower().replace("_", " ").split())
        if not heading:
            fields.append("")
            continue
        field = COLUMNS.get(heading)
        if field is None:
            raise ImportRefused(ImportFailure.UNKNOWN_COLUMN, cell.strip()[:100])
        fields.append(field)
    if not any(fields):
        raise ImportRefused(ImportFailure.NO_HEADER)
    if "first_name" not in fields:
        raise ImportRefused(ImportFailure.NO_NAME_COLUMN)
    return fields


def _csv_rows(content: bytes) -> Iterator[list[str]]:
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportRefused(ImportFailure.NOT_UTF8) from exc
    try:
        dialect = csv.Sniffer().sniff(decoded[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(decoded), dialect)
    try:
        yield from reader
    except csv.Error as exc:
        raise ImportRefused(ImportFailure.NOT_READABLE) from exc


def check_zip(content: bytes) -> None:
    """Refuse an xlsx whose zip could unpack to far more than it weighs (a "bomb")."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise ImportRefused(ImportFailure.NOT_READABLE) from exc
    if len(entries) > MAX_ZIP_ENTRIES:
        raise ImportRefused(ImportFailure.UNSAFE)
    total = 0
    for entry in entries:
        total += entry.file_size
        if total > MAX_UNPACKED_BYTES:
            raise ImportRefused(ImportFailure.UNSAFE)
        if (
            entry.file_size > MAX_RATIO * max(entry.compress_size, 1)
            and entry.file_size > RATIO_FROM_BYTES
        ):
            raise ImportRefused(ImportFailure.UNSAFE)


def _xlsx_rows(content: bytes) -> Iterator[list[str]]:
    check_zip(content)
    from openpyxl import load_workbook  # noqa: PLC0415 - only the import job needs it

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for a broken file
        raise ImportRefused(ImportFailure.NOT_READABLE) from exc
    try:
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            raise ImportRefused(ImportFailure.EMPTY)
        # Rows past the limit are never read: max_row stops openpyxl itself.
        for values in sheet.iter_rows(
            max_row=MAX_ROWS + 2, max_col=MAX_COLUMNS + 1, values_only=True
        ):
            cells = ["" if value is None else _cell(value) for value in values]
            while cells and not cells[-1]:
                cells.pop()
            yield cells
    finally:
        workbook.close()


def _cell(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel stores 1001 as 1001.0 (employee IDs, phones)
    return str(value)
