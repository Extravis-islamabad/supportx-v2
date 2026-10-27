"""CSV and XLSX exports (REQ-RPT-058) and formula-injection protection (D-100)."""

import csv
import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from app.modules.analytics.domain.tables import (
    CellKind,
    Column,
    ExportDocument,
    Sheet,
    safe_text,
)
from app.modules.analytics.infra.writers import csv_chunks, xlsx_bytes


@pytest.mark.parametrize(
    "dangerous",
    ['=HYPERLINK("http://x","y")', "+1+1", "-2+3", "@SUM(A1)", "\tcmd", "\rcmd"],
)
def test_text_that_a_spreadsheet_would_run_is_neutralised(dangerous: str) -> None:
    assert safe_text(dangerous) == f"'{dangerous}"


def test_ordinary_text_and_arabic_are_untouched() -> None:
    for text in ("Printer down", "الطابعة معطلة", "P1 \u2013 Critical", "a=b"):
        assert safe_text(text) == text


def _document(*, rtl: bool = False) -> ExportDocument:
    sheet = Sheet(
        name="Summary",
        columns=[
            Column("Priority"),
            Column("Incidents", CellKind.NUMBER),
            Column("Reopen rate", CellKind.PERCENT),
            Column("Mean time to resolve (hours)", CellKind.HOURS),
            Column("Created", CellKind.DATETIME),
        ],
        header_lines=["=Open incidents by priority", "Times in Asia/Riyadh"],
    )
    sheet.add(["=cmd|' /C calc'!A0", 8, 12.5, 7200, datetime(2026, 10, 27, 9, 30)])  # noqa: DTZ001 - wall clock of the export's zone
    sheet.add(["الأولوية العالية", 2, None, None, None])
    sheet.totals = ["Total", 10, 10.0, 7200, None]
    return ExportDocument(title="=Report", language="ar" if rtl else "en", rtl=rtl, sheets=[sheet])


def test_csv_has_a_bom_neutralised_text_and_numbers_as_numbers() -> None:
    data = b"".join(csv_chunks(_document()))
    assert data.startswith("﻿".encode())
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[0] == ["'=Open incidents by priority"]
    assert rows[3] == [
        "Priority",
        "Incidents",
        "Reopen rate",
        "Mean time to resolve (hours)",
        "Created",
    ]
    assert rows[4] == ["'=cmd|' /C calc'!A0", "8", "12.5%", "2.0", "2026-10-27 09:30"]
    assert rows[5][0] == "الأولوية العالية"
    assert rows[6][:2] == ["Total", "10"]


def test_xlsx_keeps_types_never_writes_formulas_and_is_rtl_in_arabic() -> None:
    workbook = load_workbook(io.BytesIO(xlsx_bytes(_document(rtl=True))))
    ws = workbook["Summary"]
    assert ws.sheet_view.rightToLeft is True
    values = [[c.value for c in row] for row in ws.iter_rows()]
    data = values[4]
    assert data[0] == "'=cmd|' /C calc'!A0"
    assert ws.cell(row=5, column=1).data_type == "s"  # text, not a formula
    assert data[1] == 8
    assert data[2] == pytest.approx(0.125)
    assert data[3] == pytest.approx(2.0)
    assert isinstance(data[4], datetime)
    assert values[6][:2] == ["Total", 10]
    for row in ws.iter_rows():
        for cell in row:
            assert cell.data_type != "f"
