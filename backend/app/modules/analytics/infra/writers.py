"""CSV and XLSX writers for report exports (REQ-RPT-058; D-100 D8).

Both stream: rows are appended as the database returns them, CSV straight to a file and XLSX
through openpyxl's write-only workbook (constant memory, the mode D-100 approved), so a
million-row export never sits in memory as cells. PDF is rendered in the worker image only
(infra/pdf.py, WeasyPrint).
"""

import csv
import io
from datetime import datetime
from typing import IO, TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font

from app.modules.analytics.domain.tables import (
    CellKind,
    Column,
    ExportDocument,
    Sheet,
    Value,
    safe_text,
)

if TYPE_CHECKING:
    from openpyxl.worksheet._write_only import WriteOnlyWorksheet

CSV_BOM = "﻿"  # Excel opens UTF-8 CSV with Arabic correctly only with a byte order mark
XLSX_SHEET_NAME_MAX = 31
_HOURS = 3600.0
_FORMATS = {
    CellKind.PERCENT: "0.0%",
    CellKind.HOURS: "0.00",
    CellKind.DATETIME: "yyyy-mm-dd hh:mm",
    CellKind.NUMBER: "#,##0.##",
}


def _safe(row: list[Value]) -> list[Value]:
    return [safe_text(v) if isinstance(v, str) else v for v in row]


def _plain(kind: CellKind, value: Value) -> Value:
    """A value as it goes into a file: percentages as fractions, durations as hours."""
    if value is None:
        return None
    if kind == CellKind.PERCENT and isinstance(value, int | float):
        return value / 100
    if kind == CellKind.HOURS and isinstance(value, int | float):
        return round(value / _HOURS, 2)
    return value


def _csv_text(kind: CellKind, value: Value) -> str:
    plain = _plain(kind, value)
    if plain is None:
        return ""
    if isinstance(plain, datetime):
        return plain.strftime("%Y-%m-%d %H:%M")
    if kind == CellKind.PERCENT and isinstance(plain, float):
        return f"{plain * 100:.1f}%"
    return str(plain)


class CsvSink:
    """One table as UTF-8 CSV with BOM: header lines, column headers, rows, totals."""

    def __init__(self, out: IO[bytes], columns: list[Column], header_lines: list[str]) -> None:
        self.columns = columns
        self._text = io.TextIOWrapper(out, encoding="utf-8", newline="", write_through=True)
        self._csv = csv.writer(self._text, lineterminator="\r\n")
        self._text.write(CSV_BOM)
        for line in header_lines:
            self._csv.writerow([safe_text(line)])
        if header_lines:
            self._csv.writerow([])
        self._csv.writerow([safe_text(c.header) for c in columns])

    def row(self, values: list[Value]) -> None:
        """One data row (text neutralised here)."""
        self._csv.writerow(
            [_csv_text(c.kind, v) for c, v in zip(self.columns, _safe(values), strict=True)]
        )

    def close(self) -> None:
        """Flush; the underlying file stays open for the caller."""
        self._text.flush()
        self._text.detach()


class XlsxSink:
    """A write-only workbook; sheets are filled one after another."""

    def __init__(self, *, rtl: bool) -> None:
        self.workbook = Workbook(write_only=True)
        self.rtl = rtl
        self._ws: WriteOnlyWorksheet | None = None
        self._columns: list[Column] = []

    def _cell(self, column: Column, value: Value, *, bold: bool = False) -> Cell:
        assert self._ws is not None  # noqa: S101 - sheet() is called first
        cell = WriteOnlyCell(self._ws, value=_plain(column.kind, value))
        if isinstance(value, str):
            cell.data_type = "s"  # text stays text: openpyxl would read "=..." as a formula
        elif column.kind in _FORMATS:
            cell.number_format = _FORMATS[column.kind]
        if bold:
            cell.font = Font(bold=True)
        return cell

    def sheet(self, name: str, columns: list[Column], header_lines: list[str]) -> None:
        """Start a sheet: header lines, then the bold column headers."""
        self._ws = self.workbook.create_sheet(safe_text(name)[:XLSX_SHEET_NAME_MAX] or "Report")
        self._ws.sheet_view.rightToLeft = self.rtl
        self._columns = columns
        for line in header_lines:
            self._ws.append([self._cell(Column(""), safe_text(line))])
        if header_lines:
            self._ws.append([])
        self._ws.append([self._cell(Column(""), safe_text(c.header), bold=True) for c in columns])

    def row(self, values: list[Value], *, bold: bool = False) -> None:
        """One row of the current sheet (text neutralised here)."""
        assert self._ws is not None  # noqa: S101 - sheet() is called first
        self._ws.append(
            [self._cell(c, v, bold=bold) for c, v in zip(self._columns, _safe(values), strict=True)]
        )

    def save(self, out: IO[bytes]) -> None:
        """Write the workbook."""
        self.workbook.save(out)


def csv_chunks(document: ExportDocument) -> list[bytes]:
    """The first sheet of an in-memory document as CSV (small exports and tests)."""
    buffer = io.BytesIO()
    sheet = document.sheets[0]
    sink = CsvSink(buffer, sheet.columns, sheet.header_lines)
    for row in sheet.rows:
        sink.row(row)
    if sheet.totals is not None:
        sink.row(sheet.totals)
    sink.close()
    return [buffer.getvalue()]


def xlsx_bytes(document: ExportDocument) -> bytes:
    """Every sheet of an in-memory document as one workbook (small exports and tests)."""
    sink = XlsxSink(rtl=document.rtl)
    for sheet in document.sheets:
        _fill(sink, sheet)
    buffer = io.BytesIO()
    sink.save(buffer)
    return buffer.getvalue()


def _fill(sink: XlsxSink, sheet: Sheet) -> None:
    sink.sheet(sheet.name, sheet.columns, sheet.header_lines)
    for row in sheet.rows:
        sink.row(row)
    if sheet.totals is not None:
        sink.row(sheet.totals, bold=True)
